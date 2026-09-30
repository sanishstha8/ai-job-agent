"""Export the job database to output/jobs.xlsx."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import OUTPUT_DIR, SITE_LABELS
from .store import Store

STATUS_ORDER = ["scored", "applied", "external", "failed", "skipped", "new", "duplicate", "filtered"]
COLUMNS = [
    ("Score", 8), ("Verdict", 9), ("Status", 10), ("Site", 12), ("Title", 38), ("Company", 24),
    ("Location", 22), ("Salary", 20), ("Posted", 11), ("Summary", 60), ("Missing", 45),
    ("Red flags", 30), ("Link", 14), ("Apply link (company site)", 14), ("Cover letter", 14),
    ("Applied at", 18), ("Notes", 40), ("Filter reason", 30),
]
FILLS = {
    "applied": "C6EFCE",
    "external": "FFEB9C",
    "failed": "FFC7CE",
}


def export_excel(store: Store, path: Path | None = None) -> Path:
    path = path or OUTPUT_DIR / "jobs.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    for job in store.jobs():
        details = job.score_details
        rows.append({
            "Score": job.score,
            "Verdict": job.verdict,
            "Status": job.status,
            "Site": SITE_LABELS.get(job.site, job.site),
            "Title": job.title,
            "Company": job.company,
            "Location": job.location,
            "Salary": job.salary_text(),
            "Posted": job.date_posted,
            "Summary": details.get("summary", ""),
            "Missing": "; ".join(details.get("missing_requirements", [])),
            "Red flags": "; ".join(details.get("red_flags", [])),
            "Link": job.job_url,
            "Apply link (company site)": job.job_url_direct or "",
            "Cover letter": job.cover_letter_path or "",
            "Applied at": job.applied_at or "",
            "Notes": job.apply_notes or "",
            "Filter reason": job.filter_reason or "",
        })
    df = pd.DataFrame(rows, columns=[c for c, _ in COLUMNS])
    if not df.empty:
        df["_order"] = df["Status"].map({s: i for i, s in enumerate(STATUS_ORDER)}).fillna(99)
        df = df.sort_values(["_order", "Score"], ascending=[True, False], na_position="last").drop(columns="_order")

    try:
        _write(df, path)
    except PermissionError:  # the file is open in Excel
        path = path.with_name(f"jobs_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        _write(df, path)
    return path


def _write(df: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Jobs")
        ws = writer.sheets["Jobs"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, (name, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
            ws.cell(row=1, column=i).font = Font(bold=True)
        link_cols = {name: i for i, (name, _) in enumerate(COLUMNS, start=1)
                     if name in ("Link", "Apply link (company site)", "Cover letter")}
        status_col = next(i for i, (n, _) in enumerate(COLUMNS, start=1) if n == "Status")
        for row in range(2, ws.max_row + 1):
            for name, col in link_cols.items():
                cell = ws.cell(row=row, column=col)
                if cell.value:
                    target = str(cell.value)
                    cell.hyperlink = target if target.startswith("http") else Path(target).resolve().as_uri()
                    cell.value = "open" if name != "Cover letter" else "PDF"
                    cell.font = Font(color="0563C1", underline="single")
            status = ws.cell(row=row, column=status_col).value
            if status in FILLS:
                ws.cell(row=row, column=status_col).fill = PatternFill("solid", fgColor=FILLS[status])
