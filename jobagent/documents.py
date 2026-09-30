"""Find and read the user's CV, cover letter and target-job description."""

from __future__ import annotations

from pathlib import Path

from .config import INPUTS_DIR

UPLOADABLE = (".pdf", ".docx")


def find_input(stem: str, extensions: tuple[str, ...] = (".pdf", ".docx", ".txt", ".md")) -> Path | None:
    for ext in extensions:
        path = INPUTS_DIR / f"{stem}{ext}"
        if path.exists():
            return path
    return None


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    return text.strip()


class Documents:
    """The user's inputs, loaded once per run."""

    def __init__(self) -> None:
        self.cv_path = find_input("cv")
        self.cover_letter_path = find_input("cover_letter")
        self.target_job_path = find_input("target_job", (".md", ".txt"))
        missing = [
            name
            for name, p in (
                ("inputs/cv.pdf (or .docx)", self.cv_path),
                ("inputs/cover_letter.docx (or .pdf/.txt)", self.cover_letter_path),
                ("inputs/target_job.md", self.target_job_path),
            )
            if p is None
        ]
        if missing:
            raise FileNotFoundError("Missing input files: " + ", ".join(missing))
        self.cv_text = read_text(self.cv_path)
        self.cover_letter_text = read_text(self.cover_letter_path)
        self.target_job_text = read_text(self.target_job_path)
        if len(self.cv_text) < 200:
            raise ValueError(
                f"Could only read {len(self.cv_text)} characters from {self.cv_path.name}. "
                "If it's a scanned PDF, export a text-based PDF or DOCX instead."
            )

    @property
    def cv_uploadable(self) -> bool:
        return self.cv_path is not None and self.cv_path.suffix.lower() in UPLOADABLE
